"""Reproduce conservative OOXML editing and preservation measurements.

Run with the exact, temporary Python dependencies shown in README.md. This is a
qualification probe, not an Office editing implementation or a security boundary.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import time
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree
from pptx import Presentation
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[2]
DOCX_FIXTURE = ROOT / "spikes/anydoc/corpus/docx/text.docx"
PPTX_FIXTURE = ROOT / "spikes/anydoc/corpus/pptx/pres.pptx"


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def archive_difference(original: bytes, revised: bytes) -> dict[str, object]:
    with zipfile.ZipFile(io.BytesIO(original)) as before:
        before_entries = {name: before.read(name) for name in before.namelist()}
    with zipfile.ZipFile(io.BytesIO(revised)) as after:
        after_entries = {name: after.read(name) for name in after.namelist()}
    common = before_entries.keys() & after_entries.keys()
    changed = sorted(
        name for name in common if before_entries[name] != after_entries[name]
    )
    media = sorted(name for name in common if "/media/" in name)
    return {
        "source_entry_count": len(before_entries),
        "result_entry_count": len(after_entries),
        "added_entries": sorted(after_entries.keys() - before_entries.keys()),
        "removed_entries": sorted(before_entries.keys() - after_entries.keys()),
        "changed_entries": changed,
        "unchanged_entries": len(common) - len(changed),
        "media_entries": len(media),
        "media_sha256_preserved": all(
            hashlib.sha256(before_entries[name]).digest()
            == hashlib.sha256(after_entries[name]).digest()
            for name in media
        ),
    }


def targeted_xml_edit(
    original: bytes, part: str, tag: str, old: str, new: str
) -> bytes:
    """Demonstrate package preservation for a narrow, prevalidated OOXML edit.

    This is deliberately not a general editor: a production implementation
    must validate relationships, text context, fields, signatures, and limits.
    """
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(original)) as source,
        zipfile.ZipFile(output, "w") as result,
    ):
        matches = 0
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == part:
                xml = etree.fromstring(data)
                for element in xml.iter(tag):
                    if element.text == old:
                        element.text = new
                        matches += 1
                data = etree.tostring(xml, encoding="UTF-8", xml_declaration=True)
            result.writestr(info, data)
        require(matches == 1, "Expected exactly one matching text node")
    return output.getvalue()


def docx_probe() -> dict[str, object]:
    original = DOCX_FIXTURE.read_bytes()
    document = Document(io.BytesIO(original))
    paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if any(run.text == "bold" for run in paragraph.runs)
    )
    runs_before = [
        (run.text, run.bold, run.style.name if run.style else None)
        for run in paragraph.runs
    ]
    target = next(run for run in paragraph.runs if run.text == "bold")
    original_bold = target.bold
    original_style = target.style.name if target.style else None
    target.text = "revised bold"
    buffer = io.BytesIO()
    started = time.perf_counter()
    document.save(buffer)
    elapsed = time.perf_counter() - started
    revised = buffer.getvalue()
    reopened = Document(io.BytesIO(revised))
    result = next(
        paragraph
        for paragraph in reopened.paragraphs
        if "revised bold" in paragraph.text
    )
    runs_after = [
        (run.text, run.bold, run.style.name if run.style else None)
        for run in result.runs
    ]
    require(
        [(text, bold, style) for text, bold, style in runs_before if text != "bold"]
        == [
            (text, bold, style)
            for text, bold, style in runs_after
            if text != "revised bold"
        ],
        "Unrelated runs changed",
    )
    revised_run = next(run for run in result.runs if run.text == "revised bold")
    require(revised_run.bold == original_bold, "Target run bold setting changed")
    require(
        (revised_run.style.name if revised_run.style else None) == original_style,
        "Target run style changed",
    )

    destructive = Document(io.BytesIO(original))
    destructive_paragraph = next(
        p for p in destructive.paragraphs if "Plain paragraph" in p.text
    )
    runs_before_count = len(destructive_paragraph.runs)
    explicit_run_styles_before = [
        run.style.name for run in destructive_paragraph.runs if run.style
    ]
    destructive_paragraph.text = destructive_paragraph.text.replace(
        "bold", "revised bold"
    )
    explicit_run_styles_after = [
        run.style.name for run in destructive_paragraph.runs if run.style
    ]
    surgical = targeted_xml_edit(
        original,
        "word/document.xml",
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t",
        "bold",
        "revised bold",
    )
    require(
        any(
            "revised bold" in paragraph.text
            for paragraph in Document(io.BytesIO(surgical)).paragraphs
        ),
        "Surgical DOCX result did not reopen with target text",
    )

    return {
        "fixture": str(DOCX_FIXTURE.relative_to(ROOT)),
        "source_bytes": len(original),
        "result_bytes": len(revised),
        "save_seconds": round(elapsed, 4),
        "edit": "replace one formatted run in a real paragraph",
        "run_formatting_preserved": True,
        "whole_paragraph_assignment_runs_before": runs_before_count,
        "whole_paragraph_assignment_runs_after": len(destructive_paragraph.runs),
        "whole_paragraph_assignment_explicit_styles_before": explicit_run_styles_before,
        "whole_paragraph_assignment_explicit_styles_after": explicit_run_styles_after,
        "archive": archive_difference(original, revised),
        "surgical_archive": archive_difference(original, surgical),
    }


def pptx_probe() -> dict[str, object]:
    original = PPTX_FIXTURE.read_bytes()
    presentation = Presentation(io.BytesIO(original))
    text_runs = [
        run
        for shape in presentation.slides[0].shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    ]
    target = next(run for run in text_runs if run.text == "Deck Title Slide")
    original_bold = target.font.bold
    target.text = "Revised Deck Title"
    buffer = io.BytesIO()
    started = time.perf_counter()
    presentation.save(buffer)
    elapsed = time.perf_counter() - started
    revised = buffer.getvalue()
    reopened = Presentation(io.BytesIO(revised))
    result_runs = [
        run
        for shape in reopened.slides[0].shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    ]
    require(
        any(
            run.text == "Revised Deck Title" and run.font.bold == original_bold
            for run in result_runs
        ),
        "PPTX result did not retain target text and bold setting",
    )
    surgical = targeted_xml_edit(
        original,
        "ppt/slides/slide1.xml",
        "{http://schemas.openxmlformats.org/drawingml/2006/main}t",
        "Deck Title Slide",
        "Revised Deck Title",
    )
    require(
        any(
            "Revised Deck Title" in shape.text
            for shape in Presentation(io.BytesIO(surgical)).slides[0].shapes
            if shape.has_text_frame
        ),
        "Surgical PPTX result did not reopen with target text",
    )
    return {
        "fixture": str(PPTX_FIXTURE.relative_to(ROOT)),
        "source_bytes": len(original),
        "result_bytes": len(revised),
        "save_seconds": round(elapsed, 4),
        "edit": "replace one title run on slide 1",
        "run_formatting_preserved": True,
        "archive": archive_difference(original, revised),
        "surgical_archive": archive_difference(original, surgical),
    }


def long_document_probe(fixture_dir: Path | None) -> dict[str, object]:
    document = Document()
    for number in range(2_000):
        document.add_paragraph(f"Paragraph {number}: " + "Measured text. " * 20)
    buffer = io.BytesIO()
    document.save(buffer)
    source = buffer.getvalue()
    if fixture_dir:
        (fixture_dir / "long.docx").write_bytes(source)
    started = time.perf_counter()
    reopened = Document(io.BytesIO(source))
    reopened.paragraphs[1_000].runs[0].text = "A controlled revision."
    result = io.BytesIO()
    reopened.save(result)
    elapsed = time.perf_counter() - started
    return {
        "paragraphs": 2_000,
        "source_bytes": len(source),
        "edit_and_save_seconds": round(elapsed, 4),
        "archive": archive_difference(source, result.getvalue()),
    }


def long_presentation_probe(fixture_dir: Path | None) -> dict[str, object]:
    presentation = Presentation()
    for number in range(120):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(0, 0, 7_000_000, 500_000)
        shape.text = f"Slide {number}: " + "Measured text. " * 20
    buffer = io.BytesIO()
    presentation.save(buffer)
    source = buffer.getvalue()
    if fixture_dir:
        (fixture_dir / "long.pptx").write_bytes(source)
    started = time.perf_counter()
    reopened = Presentation(io.BytesIO(source))
    reopened.slides[60].shapes[0].text_frame.paragraphs[0].runs[
        0
    ].text = "A controlled revision."
    result = io.BytesIO()
    reopened.save(result)
    elapsed = time.perf_counter() - started
    return {
        "slides": 120,
        "source_bytes": len(source),
        "edit_and_save_seconds": round(elapsed, 4),
        "archive": archive_difference(source, result.getvalue()),
    }


def long_pdf_fixture(fixture_dir: Path | None) -> dict[str, object] | None:
    if fixture_dir is None:
        return None
    source = PdfReader(ROOT / "spikes/anydoc/corpus/pdf/text.pdf")
    writer = PdfWriter()
    for _ in range(50):
        for page in source.pages:
            writer.add_page(page)
    target = fixture_dir / "long.pdf"
    with target.open("wb") as output:
        writer.write(output)
    return {"pages": len(writer.pages), "bytes": target.stat().st_size}


def explicit_break_fixture(fixture_dir: Path | None) -> None:
    if fixture_dir is None:
        return
    document = Document()
    document.add_paragraph("Before explicit page break")
    document.add_page_break()
    document.add_paragraph("After explicit page break")
    document.save(fixture_dir / "break.docx")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--fixture-dir", type=Path)
    arguments = parser.parse_args()
    if arguments.fixture_dir:
        arguments.fixture_dir.mkdir(parents=True, exist_ok=True)
    explicit_break_fixture(arguments.fixture_dir)
    print(
        json.dumps(
            {
                "docx": docx_probe(),
                "pptx": pptx_probe(),
                "long_docx": long_document_probe(arguments.fixture_dir),
                "long_pptx": long_presentation_probe(arguments.fixture_dir),
                "long_pdf": long_pdf_fixture(arguments.fixture_dir),
            },
            indent=2,
            sort_keys=True,
        )
    )


if __name__ == "__main__":
    main()
