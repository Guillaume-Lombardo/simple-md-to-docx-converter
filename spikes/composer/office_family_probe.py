"""Qualify narrow, explicit OOXML operation examples on the real corpus.

This fixture probe is not a document editor or a safe untrusted-input parser.
Every example changes one named package member and reopens or inspects the result.
"""

from __future__ import annotations

import io
import json
import zipfile
from collections.abc import Callable

from docx import Document
from docx.shared import RGBColor
from lxml import etree
from office_edit_probe import DOCX_FIXTURE, PPTX_FIXTURE, archive_difference, require
from PIL import Image
from pptx import Presentation

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"w": W, "p": P, "a": A}
PPTX_SHAPE_DELTA = 90_000
EXPECTED_SLIDES = 2


def replace_part(original: bytes, part: str, mutate: Callable[[bytes], bytes]) -> bytes:
    output = io.BytesIO()
    found = False
    with (
        zipfile.ZipFile(io.BytesIO(original)) as source,
        zipfile.ZipFile(output, "w") as result,
    ):
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == part:
                content = mutate(content)
                found = True
            result.writestr(info, content)
    require(found, f"Missing package member: {part}")
    revised = output.getvalue()
    difference = archive_difference(original, revised)
    require(
        difference["changed_entries"] == [part],
        f"Unexpected changed package members: {part}",
    )
    require(difference["added_entries"] == [], "Unexpected package member added")
    require(difference["removed_entries"] == [], "Unexpected package member removed")
    return revised


def xml_edit(content: bytes, mutate: Callable[[etree._Element], None]) -> bytes:
    root = etree.fromstring(content)
    mutate(root)
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True)


def docx_table(original: bytes) -> dict[str, object]:
    part = "word/document.xml"

    def change(root: etree._Element) -> None:
        cell_texts = root.xpath("//w:tbl[1]/w:tr[2]/w:tc[2]//w:t", namespaces=NS)
        require(
            len(cell_texts) == 1 and cell_texts[0].text == "B2",
            "Table target ambiguous",
        )
        cell_texts[0].text = "Revised B2"

    revised = replace_part(original, part, lambda data: xml_edit(data, change))
    document = Document(io.BytesIO(revised))
    require(
        document.tables[0].cell(1, 1).text == "Revised B2",
        "DOCX table cell not updated",
    )
    return {"changed_member": part, "reopened_target": True}


def docx_section(original: bytes) -> dict[str, object]:
    part = "word/document.xml"
    before = Document(io.BytesIO(original)).sections[0].bottom_margin
    expected = before + 57 * 635

    def change(root: etree._Element) -> None:
        margins = root.xpath("//w:sectPr/w:pgMar", namespaces=NS)
        require(len(margins) == 1, "Section margin target ambiguous")
        old = int(margins[0].get(f"{{{W}}}bottom"))
        margins[0].set(f"{{{W}}}bottom", str(old + 57))

    revised = replace_part(original, part, lambda data: xml_edit(data, change))
    after = Document(io.BytesIO(revised)).sections[0].bottom_margin
    require(after == expected, f"Unexpected DOCX bottom margin: {after}")
    return {"changed_member": part, "margin_delta_emu": after - before}


def docx_style(original: bytes) -> dict[str, object]:
    part = "word/styles.xml"

    def change(root: etree._Element) -> None:
        styles = root.xpath("//w:style[@w:styleId='Bold']", namespaces=NS)
        require(len(styles) == 1, "Bold style target ambiguous")
        run_properties = styles[0].find(f"{{{W}}}rPr")
        if run_properties is None:
            run_properties = etree.SubElement(styles[0], f"{{{W}}}rPr")
        color = run_properties.find(f"{{{W}}}color")
        if color is None:
            color = etree.SubElement(run_properties, f"{{{W}}}color")
        color.set(f"{{{W}}}val", "1753A6")

    revised = replace_part(original, part, lambda data: xml_edit(data, change))
    color = Document(io.BytesIO(revised)).styles["Bold"].font.color.rgb
    require(color == RGBColor(0x17, 0x53, 0xA6), "DOCX style color not updated")
    return {"changed_member": part, "reopened_target": True}


def docx_image(original: bytes) -> dict[str, object]:
    part = "word/media/image1.png"
    with zipfile.ZipFile(io.BytesIO(original)) as archive:
        source = archive.read(part)
    with Image.open(io.BytesIO(source)) as image:
        dimensions = image.size
        replacement = Image.new("RGBA", dimensions, (23, 83, 166, 255))
    output = io.BytesIO()
    replacement.save(output, format="PNG")
    revised = replace_part(original, part, lambda _: output.getvalue())
    document = Document(io.BytesIO(revised))
    require(len(document.inline_shapes) == 1, "DOCX image relationship lost")
    with (
        zipfile.ZipFile(io.BytesIO(revised)) as archive,
        Image.open(io.BytesIO(archive.read(part))) as image,
    ):
        require(image.size == dimensions, "Image dimensions changed")
    return {
        "changed_member": part,
        "dimensions": dimensions,
        "image_relationship_reopened": True,
    }


def pptx_table(original: bytes) -> dict[str, object]:
    part = "ppt/slides/slide2.xml"
    before = Presentation(io.BytesIO(original))
    table_shape = next(shape for shape in before.slides[1].shapes if shape.has_table)
    old_text = table_shape.table.cell(0, 0).text

    def change(root: etree._Element) -> None:
        text = root.xpath("//a:tbl[1]/a:tr[1]/a:tc[1]//a:t", namespaces=NS)
        require(
            len(text) == 1 and text[0].text == old_text, "PPTX table target ambiguous"
        )
        text[0].text = "Revised cell"

    revised = replace_part(original, part, lambda data: xml_edit(data, change))
    after = Presentation(io.BytesIO(revised))
    target = next(shape for shape in after.slides[1].shapes if shape.has_table)
    require(
        target.table.cell(0, 0).text == "Revised cell", "PPTX table cell not updated"
    )
    return {"changed_member": part, "reopened_target": True}


def pptx_object(original: bytes) -> dict[str, object]:
    part = "ppt/slides/slide1.xml"
    before = Presentation(io.BytesIO(original)).slides[0].shapes[0].left

    def change(root: etree._Element) -> None:
        offsets = root.xpath("//p:sp[1]/p:spPr/a:xfrm/a:off", namespaces=NS)
        require(len(offsets) == 1, "PPTX shape position ambiguous")
        offsets[0].set("x", str(int(offsets[0].get("x")) + PPTX_SHAPE_DELTA))

    revised = replace_part(original, part, lambda data: xml_edit(data, change))
    after = Presentation(io.BytesIO(revised)).slides[0].shapes[0].left
    require(after - before == PPTX_SHAPE_DELTA, "PPTX shape position not updated")
    return {"changed_member": part, "horizontal_delta_emu": after - before}


def pptx_slide_order(original: bytes) -> dict[str, object]:
    part = "ppt/presentation.xml"
    before = Presentation(io.BytesIO(original))
    first_title = before.slides[0].shapes[0].text
    second_title = before.slides[1].shapes[0].text

    def change(root: etree._Element) -> None:
        slide_lists = root.xpath("//p:sldIdLst", namespaces=NS)
        require(
            len(slide_lists) == 1 and len(slide_lists[0]) == EXPECTED_SLIDES,
            "PPTX slide order ambiguous",
        )
        first, second = slide_lists[0]
        slide_lists[0].remove(second)
        slide_lists[0].insert(0, second)
        require(slide_lists[0][1] is first, "PPTX slide order mutation failed")

    revised = replace_part(original, part, lambda data: xml_edit(data, change))
    after = Presentation(io.BytesIO(revised))
    require(
        after.slides[0].shapes[0].text == second_title, "PPTX first slide not reordered"
    )
    require(
        after.slides[1].shapes[0].text == first_title, "PPTX second slide not reordered"
    )
    return {"changed_member": part, "reopened_order": [second_title, first_title]}


def pptx_notes(original: bytes) -> dict[str, object]:
    part = "ppt/notesSlides/notesSlide1.xml"
    old = "Speaker note for the intro slide."
    new = "Revised speaker note for the intro slide."

    def change(root: etree._Element) -> None:
        texts = [element for element in root.iter(f"{{{A}}}t") if element.text == old]
        require(len(texts) == 1, "PPTX note target ambiguous")
        texts[0].text = new

    revised = replace_part(original, part, lambda data: xml_edit(data, change))
    Presentation(io.BytesIO(revised))
    with zipfile.ZipFile(io.BytesIO(revised)) as archive:
        root = etree.fromstring(archive.read(part))
    require(
        any(element.text == new for element in root.iter(f"{{{A}}}t")),
        "PPTX note XML not updated",
    )
    return {"changed_member": part, "reopened_package": True}


def main() -> None:
    docx_bytes = DOCX_FIXTURE.read_bytes()
    pptx_bytes = PPTX_FIXTURE.read_bytes()
    print(
        json.dumps(
            {
                "docx_table_cell": docx_table(docx_bytes),
                "docx_section_margin": docx_section(docx_bytes),
                "docx_style_color": docx_style(docx_bytes),
                "docx_image_bytes": docx_image(docx_bytes),
                "pptx_table_cell": pptx_table(pptx_bytes),
                "pptx_shape_position": pptx_object(pptx_bytes),
                "pptx_slide_order": pptx_slide_order(pptx_bytes),
                "pptx_note_text": pptx_notes(pptx_bytes),
            },
            indent=2,
            sort_keys=True,
        )
    )


if __name__ == "__main__":
    main()
