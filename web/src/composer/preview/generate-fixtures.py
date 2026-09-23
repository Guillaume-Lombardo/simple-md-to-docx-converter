"""Generate deterministic local stress inputs for the browser preview probe."""

from __future__ import annotations

import argparse
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from random import Random
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from lxml import etree
from PIL import Image


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "output", type=Path, help="Directory for local DOCX stress files"
    )
    args = parser.parse_args()
    output: Path = args.output
    output.mkdir(parents=True, exist_ok=True)

    long_document = Document()
    for index in range(2_000):
        long_document.add_paragraph(
            f"Paragraph {index}: measured long document content."
        )
    long_document.save(output / "composer-preview-long.docx")

    rng = Random(90)  # noqa: S311 - deterministic non-security test images
    media_document = Document()
    for index in range(30):
        image = Image.frombytes("RGB", (256, 256), rng.randbytes(256 * 256 * 3))
        stream = BytesIO()
        image.save(stream, format="PNG")
        stream.seek(0)
        media_document.add_paragraph(
            f"Raster sample {index}: image and following text."
        )
        media_document.add_picture(stream)
        media_document.add_paragraph(f"After image {index}.")
    media_document.save(output / "composer-preview-media.docx")

    source = Path(__file__).resolve().parents[4] / "spikes/anydoc/corpus/docx/text.docx"
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with (
        ZipFile(source) as input_archive,
        ZipFile(
            output / "composer-preview-structured.docx", "w", ZIP_DEFLATED
        ) as output_archive,
    ):
        for item in input_archive.infolist():
            data = input_archive.read(item.filename)
            if item.filename == "word/document.xml":
                root = etree.fromstring(data)
                body = root.find(f"{namespace}body")
                if body is None:
                    raise ValueError("The DOCX corpus has no document body.")
                parts = [
                    deepcopy(element)
                    for element in body
                    if element.tag != f"{namespace}sectPr"
                ]
                section = next(
                    (
                        deepcopy(element)
                        for element in body
                        if element.tag == f"{namespace}sectPr"
                    ),
                    None,
                )
                body.clear()
                for _ in range(12):
                    for element in parts:
                        body.append(deepcopy(element))
                if section is not None:
                    body.append(section)
                data = etree.tostring(
                    root, encoding="UTF-8", xml_declaration=True, standalone=True
                )
            output_archive.writestr(item, data)

    for name in (
        "composer-preview-long.docx",
        "composer-preview-media.docx",
        "composer-preview-structured.docx",
    ):
        path = output / name
        print(f"{path}: {path.stat().st_size} bytes")


if __name__ == "__main__":
    main()
